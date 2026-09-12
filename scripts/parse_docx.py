#!/usr/bin/env python3
"""
Simple parser to extract text from a .docx file into docs/chronic_fusion_extracted.txt
Usage:
  .\.venv\Scripts\python.exe scripts\parse_docx.py [path/to/docx]
If no path provided, defaults to 'chronic-fusion.docx' in the repo root.
"""
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("python-docx is not installed. Install it with:\n  .\.venv\\Scripts\\python.exe -m pip install python-docx")
    sys.exit(2)

def extract(docx_path: Path, out_path: Path):
    doc = Document(docx_path)
    lines = []

    # paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)

    # tables
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                lines.append(row_text)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return lines

if __name__ == '__main__':
    if len(sys.argv) > 1:
        in_path = Path(sys.argv[1])
    else:
        in_path = Path('chronic-fusion.docx')

    if not in_path.exists():
        print(f"Doc file not found: {in_path}\nPlease place the Word file in the repository root or pass its path as argument.")
        sys.exit(1)

    out_file = Path('docs/chronic_fusion_extracted.txt')
    try:
        lines = extract(in_path, out_file)
        print(f"Extracted {len(lines)} text lines to {out_file}")
        preview = out_file.read_text(encoding='utf-8')[:2000]
        print('\n--- Preview (first 2000 chars) ---\n')
        print(preview)
    except Exception as e:
        print('Error extracting docx:', e)
        sys.exit(1)
